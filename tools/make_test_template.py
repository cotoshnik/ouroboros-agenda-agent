#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Генерирует тестовый PPTX-шаблон templates/report_template.pptx.

Стенд-in для разработки и тестов: настоящий production-шаблон предоставляет
пользователь (путь через template_path или AGENDA_PPTX_TEMPLATE). Этот файл
содержит титульный слайд с TITLE+SUBTITLE, логотип на слайде и наследует
мастер/макеты/тему python-pptx — достаточно для проверки маппинга.
"""

import base64
import os
import tempfile

from pptx import Presentation
from pptx.util import Cm

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE_DIR, "templates", "report_template.pptx")

# 2x2 px PNG (заглушка логотипа)
LOGO_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAYAAABytg0kAAAAEklEQVR4nGP8z8DwnwEKmBgQAAA9"
    "9wED6uLzgwAAAABJRU5ErkJggg==")


def main():
    prs = Presentation()  # база шаблона: стандартный master/тема/макеты
    title_layout = prs.slide_layouts[0]  # Title Slide: TITLE + SUBTITLE
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = "Шаблон справки — заголовок"
    subtitle = slide.placeholders[1]
    subtitle.text = "дата — статус"

    # логотип на титульном слайде шаблона
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        f.write(LOGO_PNG)
        logo_path = f.name
    slide.shapes.add_picture(logo_path, Cm(1), Cm(1), width=Cm(2.5))
    os.unlink(logo_path)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print("Шаблон создан:", OUT)


if __name__ == "__main__":
    main()
