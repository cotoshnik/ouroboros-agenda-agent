#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Тесты MCP-сервера agenda_briefing (unit + integration + валидация файлов).

Запуск:  python3 -m unittest discover -s tests -v   (из корня проекта)
Данные тестов пишутся во временный каталог (AGENDA_DATA_DIR до импорта
сервера), рабочая папка data/ не затрагивается. Значения в тестах —
тестовые фикстуры, не production-данные.
"""

import json
import os
import shutil
import sys
import tempfile
import threading
import unittest
import urllib.request
import zipfile

TEST_DATA_DIR = tempfile.mkdtemp(prefix="agenda_test_")
os.environ["AGENDA_DATA_DIR"] = TEST_DATA_DIR

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "mcp"))
import server  # noqa: E402
import pptx_export  # noqa: E402

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(BASE_DIR, "templates", "report_template.pptx")

DATE = "2026-10-26"          # базовая дата тестов
DATE2 = "2026-10-25"         # для изоляции отдельных тестов


def ensure_template():
    if not os.path.exists(TEMPLATE):
        sys.path.insert(0, os.path.join(BASE_DIR, "tools"))
        import make_test_template
        make_test_template.main()


def save(iid, value, unit="", source="тест", date=DATE):
    return server.tool_save_observation(
        {"indicator_id": iid, "date": date, "value": value, "unit": unit, "source": source})


class TestDataLayer(unittest.TestCase):
    def test_catalog_reads(self):
        catalog = server.load_catalog()
        self.assertEqual(catalog["total"], 67)
        self.assertTrue(any(i["base_block"] for i in catalog["indicators"]))

    def test_save_and_get_observation(self):
        save(1, 3.77, "млн чел.", date="2026-11-01")
        obs = server.tool_get_observations({"date": "2026-11-01"})
        self.assertEqual(obs["total"], 1)
        self.assertEqual(obs["observations"][0]["value"], 3.77)

    def test_unknown_indicator_rejected(self):
        with self.assertRaises(server.ToolError):
            save(999, 1, date="2026-11-02")

    def test_invalid_date_format_rejected(self):
        with self.assertRaises(server.ToolError):
            save(1, 1, date="26.10.2026")

    def test_nonexistent_calendar_date_rejected(self):
        with self.assertRaises(server.ToolError):
            save(1, 1, date="2026-02-30")

    def test_bad_types_rejected(self):
        with self.assertRaises(server.ToolError):
            save("26", 1, date="2026-11-02")   # id строкой
        with self.assertRaises(server.ToolError):
            save(1, True, date="2026-11-02")   # bool как значение
        with self.assertRaises(server.ToolError):
            save(1, "x" * 600, date="2026-11-02")  # слишком длинная строка

    def test_completeness(self):
        save(1, 3.77, "млн чел.", date="2026-11-03")
        result = server.tool_check_completeness({"date": "2026-11-03"})
        self.assertFalse(result["complete"])
        self.assertEqual(result["base_total"], 49)
        self.assertEqual(result["collected_base"], 1)

    def test_hot_rules(self):
        save(26, 99.41, "%", date="2026-11-04")   # ниже порога 99.5
        save(28, 0, "шт.", date="2026-11-04")     # 0 — не горячее
        hot = server.tool_detect_hot_topics({"date": "2026-11-04"})["hot_topics"]
        ids = [h["indicator_id"] for h in hot]
        self.assertIn(26, ids)
        self.assertNotIn(28, ids)

    def test_zero_is_not_meaningless(self):
        # нулевое значение валидно и доступно для отбора
        save(33, 0, "шт.", date="2026-11-05")
        obs = server._obs_map("2026-11-05")
        self.assertEqual(obs[33]["value"], 0)

    def test_missing_value_marks_no_data(self):
        draft = server.tool_build_report_draft({"date": "2026-11-06"})["draft_markdown"]
        self.assertIn(server.NO_DATA, draft)
        self.assertNotIn("Мониторинг открытых источников", draft)

    def test_corrupted_observations_not_clobbered(self):
        path = server.observations_path("2026-10-24")
        with open(path, "w", encoding="utf-8") as f:
            f.write("{not json")
        with self.assertRaises(server.ToolError):
            server.tool_get_observations({"date": "2026-10-24"})
        with self.assertRaises(server.ToolError):
            save(1, 1, date="2026-10-24")
        with open(path, encoding="utf-8") as f:
            self.assertEqual(f.read(), "{not json")  # файл не перезаписан


class TestSelection(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        save(26, 99.41, "%", date=DATE)        # hot
        save(58, 3, "атаки", date=DATE)        # hot
        save(1, 3.77, "млн чел.", date=DATE)   # не hot
        save(15, 1.9, "млрд руб.", date=DATE)  # dynamic, не hot

    def test_proposal_includes_only_hot(self):
        result = server.tool_select_presentation_indicators({"date": DATE})
        self.assertEqual(result["mode"], "proposal")
        included = [d["indicator_id"] for d in result["decisions"] if d["include"]]
        self.assertEqual(sorted(included), [26, 58])

    def test_server_validation_rejects_bad_ids(self):
        selections = [
            {"indicator_id": 1, "include": True, "priority": "normal",
             "justification": "ежедневный контроль клиентопотока", "ref": "observation"},
            {"indicator_id": 999, "include": True, "priority": "high",
             "justification": "нет такого", "ref": "observation"},
            {"indicator_id": 8, "include": True, "priority": "high",
             "justification": "нет наблюдения за дату", "ref": "observation"},
            {"indicator_id": 15, "include": True, "priority": "weird",
             "justification": "динамический, есть отклонение", "ref": "observation"},
        ]
        result = server.tool_select_presentation_indicators(
            {"date": DATE, "selections": selections})
        ids = result["selected_indicator_ids"]
        self.assertIn(1, ids)
        self.assertIn(15, ids)                 # priority нормализован
        self.assertIn(26, ids)                 # hot принудительно добавлен сервером
        self.assertIn(58, ids)
        self.assertNotIn(999, ids)
        self.assertNotIn(8, ids)               # нет фактического наблюдения
        self.assertTrue(result["warnings"])

    def test_audit_written(self):
        server.tool_select_presentation_indicators(
            {"date": DATE, "selections": [
                {"indicator_id": 1, "include": True, "priority": "normal",
                 "justification": "контроль", "ref": "observation"}]})
        audit = server._read_audit(DATE)
        self.assertEqual(audit["source"], "analyst")
        self.assertTrue(audit["decisions"])

    def test_selection_deterministic(self):
        selections = [{"indicator_id": 1, "include": True, "priority": "normal",
                       "justification": "контроль", "ref": "observation"}]
        r1 = server.tool_select_presentation_indicators({"date": DATE, "selections": selections})
        r2 = server.tool_select_presentation_indicators({"date": DATE, "selections": selections})
        self.assertEqual(r1["selected_indicator_ids"], r2["selected_indicator_ids"])


class TestDocxExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        save(3, 1234567890123, "млн чел. «тест» — ☃", date="2026-10-23")  # unicode + большое число + длинное название

    def test_docx_values_match_observations(self):
        result = server.tool_export_report_docx({"date": "2026-10-23"})
        path = result["docx_path"]
        self.assertTrue(zipfile.is_zipfile(path))
        with zipfile.ZipFile(path) as z:
            document = z.read("word/document.xml").decode("utf-8")
        self.assertIn("1234567890123", document)
        self.assertIn("☃", document)
        self.assertIn(server.NO_DATA, document)

    def test_docx_atomic_no_tmp_left(self):
        server.tool_export_report_docx({"date": DATE})
        leftovers = [f for f in os.listdir(server.DATA_DIR) if ".tmp" in f]
        self.assertEqual(leftovers, [])


class TestPptxExport(unittest.TestCase):
    PDATE = "2026-11-11"

    @classmethod
    def setUpClass(cls):
        ensure_template()
        save(26, 99.41, "%", date=cls.PDATE)      # hot
        save(58, 3, "атаки", date=cls.PDATE)      # hot
        save(1, 3.77, "млн чел.", date=cls.PDATE)
        save(15, 1.9, "млрд руб.", date=cls.PDATE)

    def test_template_required_without_template(self):
        result = server.tool_export_report_pptx({"date": self.PDATE})
        self.assertEqual(result["status"], "template_required")
        self.assertIsNone(result["pptx_path"])

    def test_corrupted_template_rejected(self):
        bad = os.path.join(TEST_DATA_DIR, "bad.pptx")
        with open(bad, "wb") as f:
            f.write(b"not a zip")
        with self.assertRaises(server.ToolError):
            server.tool_export_report_pptx({"date": self.PDATE, "template_path": bad})

    def test_missing_template_rejected(self):
        with self.assertRaises(server.ToolError):
            server.tool_export_report_pptx(
                {"date": self.PDATE, "template_path": "/nonexistent/x.pptx"})

    def test_inspect_template(self):
        info = server.tool_inspect_pptx_template({"template_path": TEMPLATE})
        self.assertGreaterEqual(len(info["layouts"]), 1)
        self.assertGreaterEqual(info["slide_count"], 1)

    def test_pptx_content_and_template_preserved(self):
        selection = [{"indicator_id": 1, "include": True, "priority": "normal",
                      "justification": "ежедневный контроль", "ref": "observation"}]
        result = server.tool_export_report_pptx(
            {"date": self.PDATE, "template_path": TEMPLATE, "selection": selection})
        self.assertEqual(result["status"], "success")
        self.assertEqual(sorted(result["selected_indicator_ids"]), [1, 26, 58])

        path = result["pptx_path"]
        self.assertTrue(zipfile.is_zipfile(path))
        from pptx import Presentation
        prs_out = Presentation(path)
        prs_tpl = Presentation(TEMPLATE)
        # тема/макеты шаблона сохранены
        self.assertEqual(len(prs_out.slide_layouts), len(prs_tpl.slide_layouts))
        # логотип титульного слайда шаблона сохранён
        pics = sum(1 for sh in prs_out.slides[0].shapes if sh.shape_type == 13)
        self.assertGreaterEqual(pics, 1)
        # значения совпадают с observations, нет новостей и лишних показателей
        all_text = "\n".join(sh.text_frame.text for s in prs_out.slides
                             for sh in s.shapes if sh.has_text_frame)
        self.assertIn("99.41", all_text)
        self.assertIn("3.77", all_text)
        self.assertNotIn("1.9", all_text)              # id=15 не выбран — не выносится
        self.assertNotIn("Мониторинг открытых источников", all_text)
        self.assertNotIn("ria.ru", all_text)
        self.assertIn("ЧЕРНОВИК", all_text)
        # нет пустых плейсхолдеров на созданных слайдах
        for s in prs_out.slides:
            for ph in s.placeholders:
                self.assertTrue(ph.text_frame.text.strip(),
                                "пустой плейсхолдер на слайде")

    def test_pptx_repeat_deterministic(self):
        selection = [{"indicator_id": 1, "include": True, "priority": "normal",
                      "justification": "контроль", "ref": "observation"}]
        r1 = server.tool_export_report_pptx({"date": self.PDATE, "template_path": TEMPLATE,
                                             "selection": selection})
        from pptx import Presentation
        t1 = [sh.text_frame.text for s in Presentation(r1["pptx_path"]).slides
              for sh in s.shapes if sh.has_text_frame]
        r2 = server.tool_export_report_pptx({"date": self.PDATE, "template_path": TEMPLATE,
                                             "selection": selection})
        t2 = [sh.text_frame.text for s in Presentation(r2["pptx_path"]).slides
              for sh in s.shapes if sh.has_text_frame]
        self.assertEqual(r1["selected_indicator_ids"], r2["selected_indicator_ids"])
        self.assertEqual(t1, t2)


class TestExportContract(unittest.TestCase):
    CDATE = "2026-11-12"

    @classmethod
    def setUpClass(cls):
        ensure_template()
        save(26, 99.41, "%", date=cls.CDATE)      # hot
        save(1, 3.77, "млн чел.", date=cls.CDATE)

    def test_success_contract(self):
        result = server.tool_export_report({"date": self.CDATE, "template_path": TEMPLATE})
        self.assertEqual(result["status"], "success")
        self.assertTrue(os.path.getsize(result["docx_path"]) > 0)
        self.assertTrue(os.path.getsize(result["pptx_path"]) > 0)
        self.assertIn(26, result["selected_indicator_ids"])
        self.assertGreater(result["missing_base_count"], 0)

    def test_one_file_failure_keeps_other(self):
        result = server.tool_export_report(
            {"date": self.CDATE, "template_path": "/nonexistent/x.pptx"})
        self.assertEqual(result["status"], "error")
        self.assertEqual(result["failed_stage"], "pptx")
        self.assertIsNone(result["pptx_path"])
        self.assertTrue(os.path.exists(result["docx_path"]))  # DOCX не удалён

    def test_export_without_observations(self):
        result = server.tool_export_report({"date": "2026-10-22", "template_path": TEMPLATE})
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["selected_indicator_ids"], [])
        self.assertEqual(result["missing_base_count"], 49)


class TestHttpIntegration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ensure_template()
        from http.server import ThreadingHTTPServer
        cls.httpd = ThreadingHTTPServer(("127.0.0.1", 0), server.MCPHandler)
        cls.httpd.daemon_threads = True
        cls.port = cls.httpd.server_address[1]
        cls.thread = threading.Thread(target=cls.httpd.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()
        cls.httpd.server_close()

    def call(self, payload, raw=None):
        data = raw.encode("utf-8") if raw is not None else json.dumps(payload).encode("utf-8")
        req = urllib.request.Request(
            "http://127.0.0.1:%d/mcp" % self.port, data=data,
            headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=10) as resp:
                return resp.status, json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read().decode("utf-8"))

    def test_tools_list_has_no_news_tools(self):
        status, resp = self.call({"jsonrpc": "2.0", "id": 1, "method": "tools/list"})
        names = [t["name"] for t in resp["result"]["tools"]]
        self.assertNotIn("fetch_web_sources", names)
        self.assertNotIn("get_signals", names)
        self.assertEqual(len(names), 12)

    def test_initialize_ping_unknown(self):
        status, resp = self.call({"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {}})
        self.assertEqual(resp["result"]["serverInfo"]["name"], "agenda_briefing")
        _, resp = self.call({"jsonrpc": "2.0", "id": 2, "method": "ping"})
        self.assertEqual(resp["result"], {})
        _, resp = self.call({"jsonrpc": "2.0", "id": 3, "method": "no/such"})
        self.assertEqual(resp["error"]["code"], -32601)

    def test_malformed_json(self):
        status, resp = self.call(None, raw="{broken")
        self.assertEqual(status, 400)

    def test_oversize_body_rejected(self):
        data = b'{"jsonrpc":"2.0","id":1,"method":"ping","pad":"' + b"x" * (1024 * 1024 + 10) + b'"}'
        req = urllib.request.Request("http://127.0.0.1:%d/mcp" % self.port, data=data)
        try:
            urllib.request.urlopen(req, timeout=10)
            self.fail("oversize запрос должен быть отклонён")
        except urllib.error.HTTPError as e:
            self.assertEqual(e.code, 413)
        except urllib.error.URLError:
            pass  # сервер закрыл соединение до конца передачи — тоже отклонение

    def test_path_traversal_date_rejected(self):
        _, resp = self.call({"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                             "params": {"name": "get_observations",
                                        "arguments": {"date": "../../etc"}}})
        self.assertTrue(resp["result"]["isError"])

    def test_concurrent_saves_no_lost_updates(self):
        date = "2026-10-21"
        ids = [1, 2, 3, 8, 9, 10, 11, 12, 13, 26, 27, 28, 30, 32, 33, 34, 55, 56, 57, 58]

        def worker(iid):
            self.call({"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                       "params": {"name": "save_observation",
                                  "arguments": {"indicator_id": iid, "date": date,
                                                "value": iid, "source": "конкурентный тест"}}})

        threads = [threading.Thread(target=worker, args=(i,)) for i in ids]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        _, resp = self.call({"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                             "params": {"name": "get_observations", "arguments": {"date": date}}})
        got = json.loads(resp["result"]["content"][0]["text"])
        self.assertEqual(got["total"], len(ids), "потеряны обновления при гонке")

    def test_export_during_updates(self):
        date = "2026-10-21"
        errors = []

        def exporter():
            try:
                self.call({"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                           "params": {"name": "export_report",
                                      "arguments": {"date": date, "template_path": TEMPLATE}}})
            except Exception as e:
                errors.append(e)

        def saver():
            for i in range(20):
                self.call({"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                           "params": {"name": "save_observation",
                                      "arguments": {"indicator_id": 1, "date": date,
                                                    "value": i, "source": "гонка"}}})

        t1, t2 = threading.Thread(target=exporter), threading.Thread(target=saver)
        t1.start(); t2.start(); t1.join(); t2.join()
        self.assertEqual(errors, [])

    def test_full_cycle_e2e(self):
        date = "2026-10-20"
        for iid, val, unit in [(26, 99.2, "%"), (56, 2, "шт."), (1, 4.01, "млн чел.")]:
            self.call({"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                       "params": {"name": "save_observation",
                                  "arguments": {"indicator_id": iid, "date": date,
                                                "value": val, "unit": unit,
                                                "source": "e2e-фикстура"}}})
        _, resp = self.call({"jsonrpc": "2.0", "id": 2, "method": "tools/call",
                             "params": {"name": "select_presentation_indicators",
                                        "arguments": {"date": date, "selections": [
                                            {"indicator_id": 1, "include": True,
                                             "priority": "normal",
                                             "justification": "ежедневный контроль",
                                             "ref": "observation"}]}}})
        selection = json.loads(resp["result"]["content"][0]["text"])
        self.assertIn(26, selection["selected_indicator_ids"])
        _, resp = self.call({"jsonrpc": "2.0", "id": 3, "method": "tools/call",
                             "params": {"name": "export_report",
                                        "arguments": {"date": date, "template_path": TEMPLATE}}})
        result = json.loads(resp["result"]["content"][0]["text"])
        self.assertEqual(result["status"], "success")
        self.assertTrue(os.path.getsize(result["docx_path"]) > 0)
        self.assertTrue(os.path.getsize(result["pptx_path"]) > 0)
        audit = server._read_audit(date)
        self.assertIsNotNone(audit)


def tearDownModule():
    shutil.rmtree(TEST_DATA_DIR, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
